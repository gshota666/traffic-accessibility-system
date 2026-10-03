from services.research import research_report,export_research
import time
import asyncio
import json
import logging
import secrets
import shutil
import uuid
import csv
import io
import zipfile
import re
from contextlib import asynccontextmanager
from datetime import datetime,date as Date,time as dtime,timezone,timedelta
from pathlib import Path
from typing import Literal
from urllib.parse import urlparse
from fastapi import FastAPI,UploadFile,File,HTTPException,Request,Query
from fastapi.exceptions import RequestValidationError
from fastapi.responses import FileResponse,JSONResponse,RedirectResponse,StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel,Field
from backend.db import Database
from services.capacity import estimate_capacity
from services.settings import Settings
from services.worker import Worker,now
from services.importer import headers,ingest,read_rows,canonical_crs
from services.exporter import export_details,export_summary,summaries,export_package,export_paper,export_quality,DATA_USE_LABELS
from services.research_plan import FINAL_TEMPLATE_VERSION,is_final_plan
from utils.version import APP_VERSION
from utils.paths import data_dir,resource
from utils.geo import convert,SYSTEMS

class Mapping(BaseModel):
    mapping:dict[str,int|None]
    crs:Literal['WGS84','GCJ-02','BD-09']='WGS84'
class JobInput(BaseModel):
    name:str=Field(min_length=1,max_length=120)
    estimate_seconds:float=Field(default=0.75,ge=0.01,le=300,allow_inf_nan=False)
    acknowledge_capacity:bool=False
    dataset_id:str
    # Optional subset of validated source point primary keys.  An empty list
    # keeps the legacy behaviour (use every valid source row), while the UI
    # can use this field for a small, no-copy trial drawn directly from the
    # formal source CSV.
    source_point_ids:list[int]=Field(default_factory=list,max_length=10000)
    target_ids:list[int]=Field(min_length=1,max_length=1000)
    dates:list[Date]=Field(default_factory=list,max_length=31)
    day_type:Literal['workday','weekend','holiday','custom']='custom'
    query_order:Literal['randomized','original']='randomized'
    random_seed:int|None=Field(default=None,ge=0,le=2147483647)
    date:Date|None=None
    times:list[str]=Field(default_factory=list,max_length=24)
    mode:Literal['exact','saving']='exact'
    candidates:int=Field(default=5,ge=1,le=1000)
    provider:Literal['mock','amap']='mock'
    strategy:Literal['32','33','34','35','36','37','38','39','40','41','42','43','44','45']='32'
    collection_mode:Literal['immediate','scheduled']='immediate'
    window_minutes:int=Field(default=30,ge=1,le=120)
    acknowledge_estimate:bool=False
    directions:list[Literal['outbound','inbound']]=Field(default_factory=lambda:['outbound'],min_length=1,max_length=2)
    study_id:str=Field(default='',max_length=120)
    # Optional per-date schedule.  The legacy dates/times fields remain
    # accepted so old clients and saved drafts continue to work.
    schedule:list['ScheduleInput']=Field(default_factory=list,max_length=31)
    # Versioned research calendar metadata.  These fields are optional so
    # saved jobs created by older clients continue to load unchanged.
    schedule_template:str=Field(default='',max_length=120)
    schedule_adjusted:bool=False
    # The final Word fixes the clock time but deliberately leaves the
    # allowed request window for user confirmation.  Keep that confirmation
    # separate from the provisional per-row window value.
    schedule_window_confirmed:bool=False
    # Purpose is explicit metadata, not an inference from provider or sample
    # size.  Existing tasks have no value and are exported as “未标注”.
    data_use:Literal['unmarked','mock','pilot','formal']='unmarked'

class ScheduleInput(BaseModel):
    date:Date
    phase:str=Field(default='',max_length=120)
    day_type:Literal['workday','weekend','holiday','custom']='custom'
    times:list[str]=Field(default_factory=list,max_length=24)
    window_minutes:int=Field(default=20,ge=1,le=120)
    notes:str=Field(default='',max_length=500)
    # Optional per-date direction override.  When omitted, the task-level
    # directions are inherited; this keeps old saved tasks compatible while
    # allowing a research calendar to use different directions on different
    # dates.
    directions:list[Literal['outbound','inbound']]|None=Field(default=None,max_length=2)

class TargetInput(BaseModel):
    name:str=Field(min_length=1,max_length=120)
    lon:float=Field(ge=-180,le=180,allow_inf_nan=False)
    lat:float=Field(ge=-90,le=90,allow_inf_nan=False)
    crs:Literal['WGS84','GCJ-02','BD-09']='GCJ-02'
    note:str=Field(default='',max_length=500)
    target_id:str=Field(default='',max_length=120)
    province:str=Field(default='',max_length=80)
    city:str=Field(default='',max_length=80)
    county:str=Field(default='',max_length=80)
    village_name:str=Field(default='',max_length=160)
    entrance_name:str=Field(default='',max_length=160)
    coordinate_source:str=Field(default='',max_length=200)
    verification_status:str=Field(default='',max_length=80)
class TargetCrsInput(BaseModel):
    crs:Literal['WGS84','GCJ-02','BD-09']
class TargetIdsInput(BaseModel):
    ids:list[int]=Field(min_length=1,max_length=10000)
class JobIdsInput(BaseModel):
    ids:list[str]=Field(min_length=1,max_length=200)
class DatasetIdsInput(BaseModel):
    ids:list[str]=Field(min_length=1,max_length=200)
class SettingInput(BaseModel):
    provider:Literal['mock','amap']='mock'
    qps:float=Field(default=3,gt=0,le=20,allow_inf_nan=False)
    max_concurrency:int=Field(default=3,ge=1,le=8)
    daily_limit:int=Field(default=1000,ge=1,le=1000000)
    retries:int=Field(default=3,ge=0,le=5)
    strategy:str='32'
    api_key:str=Field(default='',max_length=256)
    plaintext:bool=False
class Rename(BaseModel):
    name:str|None=Field(default=None,min_length=1,max_length=120)
    data_use:Literal['unmarked','mock','pilot','formal']|None=None

def create_app(root=None,token=None,run_worker=True,on_shutdown=None):
    root=data_dir(root); db=Database(root); settings=Settings(db,root); worker=Worker(db,settings)
    token=token or secrets.token_urlsafe(32)
    @asynccontextmanager
    async def lifespan(app):
        if run_worker: worker.start()
        yield
        await asyncio.to_thread(worker.close)
    app=FastAPI(title='交通可达性批量分析系统',lifespan=lifespan,docs_url=None,redoc_url=None,openapi_url=None)
    app.state.db=db; app.state.worker=worker; app.state.token=token
    @app.middleware('http')
    async def local_only(request,call_next):
        host=request.headers.get('host','').split(':')[0]
        if host not in {'127.0.0.1','localhost','testserver'}:
            return JSONResponse({'detail':'仅允许本机访问'},403)
        if request.url.path=='/health': return JSONResponse({'app':'TrafficAccessibility','ready':True})
        origin=request.headers.get('origin')
        if origin and origin!=str(request.base_url).rstrip('/'):
            return JSONResponse({'detail':'拒绝跨站请求'},403)
        if request.url.path=='/' and secrets.compare_digest(request.query_params.get('token',''),token):
            response=RedirectResponse('/',status_code=303)
            response.set_cookie('traffic_session',token,httponly=True,samesite='strict')
        elif not secrets.compare_digest(request.cookies.get('traffic_session',request.headers.get('x-session-token','')),token):
            return JSONResponse({'detail':'请从应用入口重新打开浏览器'},401)
        else:
            response=await call_next(request)
        response.headers['Cache-Control']='no-store'
        response.headers['Referrer-Policy']='no-referrer'
        response.headers['X-Content-Type-Options']='nosniff'
        response.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
        return response
    @app.exception_handler(RequestValidationError)
    async def invalid_request(request,exc):
        labels={'lon':'经度','lat':'纬度','name':'名称','note':'备注','qps':'QPS','daily_limit':'每日请求预算'}
        messages=[]
        for error in exc.errors():
            field=str(error['loc'][-1]); label=labels.get(field,'输入字段')
            if field in {'lon','lat'}:
                bound=180 if field=='lon' else 90
                message=f'{label}必须是 -{bound} 到 {bound} 之间的数字，请检查小数点。'
            elif error['type']=='string_too_short': message=f'{label}不能为空。'
            elif error['type']=='string_too_long': message=f'{label}过长，请缩短后重试。'
            else: message=f'{label}格式或取值不正确，请检查后重试。'
            if message not in messages: messages.append(message)
        return JSONResponse({'detail':' '.join(messages)},422)
    @app.exception_handler(ValueError)
    async def bad_value(request,exc): return JSONResponse({'detail':str(exc)},400)
    @app.get('/')
    def index(): return FileResponse(resource('frontend/index.html'))
    app.mount('/static',StaticFiles(directory=resource('frontend')),name='static')
    def get_job(jid):
        job=db.one('SELECT * FROM jobs WHERE id=?',(jid,))
        if not job: raise HTTPException(404,'任务不存在')
        job['config']=json.loads(job['config'])
        # Legacy tasks predate the explicit purpose marker.  This is a
        # presentation default only; the historical config and observations
        # are left untouched until the user explicitly saves a marker.
        job['config'].setdefault('data_use','unmarked')
        job['collection']=db.one("SELECT min(requested_at) started_at,max(CASE WHEN requested_at IS NOT NULL THEN coalesce(fetched_at,calculated_at) END) finished_at,count(requested_at) requested_count FROM results WHERE job_id=?",(jid,))
        job['performance']=db.one("""SELECT count(*) attempts,
            avg(rate_wait_seconds) avg_rate_wait_seconds,
            avg(provider_duration_seconds) avg_provider_seconds,
            avg(raw_persistence_duration_seconds) avg_raw_persistence_seconds,
            avg(segment_persistence_duration_seconds) avg_segment_persistence_seconds,
            avg(database_duration_seconds) avg_database_seconds,
            avg(persistence_duration_seconds) avg_persistence_seconds,
            avg(total_duration_seconds) avg_total_seconds,
            max(total_duration_seconds) max_total_seconds,
            max(max_concurrency) max_concurrency
            FROM request_attempts WHERE job_id=?""",(jid,))
        return job
    @app.get('/api/version')
    def version(): return {'version':APP_VERSION}
    @app.get('/api/dashboard')
    def dashboard():
        result={}
        for key,table in [('jobs','jobs'),('points','points')]: result[key]=db.one(f'SELECT count(*) n FROM {table}')['n']
        result['routes']=db.one("SELECT count(*) n FROM results WHERE status='success'")['n']
        result['calls']=db.one('SELECT coalesce(sum(calls),0) n FROM usage')['n']
        return result
    @app.get('/api/settings')
    def get_settings(): return settings.public()
    @app.put('/api/settings')
    def save_settings(body:SettingInput):
        settings.save(body.model_dump(exclude={'api_key','plaintext'}),body.api_key,body.plaintext)
        return settings.public()
    @app.get('/api/usage')
    def usage(): return db.rows('SELECT * FROM usage ORDER BY day DESC LIMIT 90')
    @app.get('/api/events')
    def events(job_id:str=''):
        return db.rows('SELECT * FROM events WHERE (?=\'\' OR job_id=?) ORDER BY id DESC LIMIT 100',(job_id,job_id))
    @app.post('/api/imports')
    def upload(file:UploadFile=File(...)):
        suffix=Path(file.filename or '').suffix.lower()
        if suffix not in {'.csv','.xlsx'}: raise HTTPException(400,'仅支持 CSV 或 XLSX')
        ident=uuid.uuid4().hex; path=root/'uploads'/f'{ident}{suffix}'
        size=0
        try:
            with path.open('wb') as dest:
                while chunk:=file.file.read(1024*1024):
                    size+=len(chunk)
                    if size>100*1024*1024: raise ValueError('文件超过 100 MB 限制，请拆分导入')
                    dest.write(chunk)
            cols,mapping=headers(path)
            # If the file contains a declared coordinate-system column, use
            # its first non-empty valid value as a helpful default in the
            # import dialog.  The row-level value remains authoritative.
            detected_crs=''
            ci=mapping.get('coordinate_system')
            if ci is not None:
                rows=read_rows(path); next(rows,None)
                for values in rows:
                    if ci < len(values) and str(values[ci] or '').strip():
                        detected_crs=canonical_crs(values[ci]) or ''
                        if detected_crs: break
                rows.close()
        except Exception as exc:
            path.unlink(missing_ok=True)
            raise HTTPException(400,str(exc) if isinstance(exc,ValueError) else 'Excel/CSV 文件损坏或格式无效') from None
        db.execute('INSERT INTO datasets VALUES(?,?,?,?,?,?)',(ident,file.filename,now(),json.dumps(cols),'',json.dumps({'uploaded_file':path.name})))
        return dict(id=ident,columns=cols,mapping=mapping,crs=detected_crs)
    @app.post('/api/imports/{ident}/validate')
    def validate_import(ident:str,body:Mapping):
        if not db.one('SELECT id FROM datasets WHERE id=?',(ident,)): raise HTTPException(404,'数据集不存在')
        if db.one('SELECT id FROM jobs WHERE dataset_id=?',(ident,)): raise HTTPException(409,'该数据集已用于任务，不可修改映射')
        files=list((root/'uploads').glob(ident+'.*'))
        if len(files)!=1: raise HTTPException(404,'导入原文件不存在')
        try: return ingest(db,ident,files[0],body.mapping,body.crs)
        except (ValueError,TypeError): raise
        except Exception: raise HTTPException(400,'读取数据失败：请检查工作簿内容') from None
    @app.get('/api/datasets')
    def datasets():
        rows=db.rows('''SELECT d.*,
            (SELECT count(*) FROM jobs j WHERE j.dataset_id=d.id) AS job_count
            FROM datasets d ORDER BY d.created DESC''')
        for r in rows: r['stats']=json.loads(r['stats'])
        return rows
    @app.post('/api/datasets/bulk-delete')
    def bulk_delete_datasets(body:DatasetIdsInput):
        ids=list(dict.fromkeys(body.ids))
        placeholders=','.join('?' for _ in ids)
        rows=db.rows(f'SELECT id,stats FROM datasets WHERE id IN ({placeholders})',ids)
        if len(rows)!=len(ids): raise HTTPException(404,'部分数据已不存在，请刷新页面后重试')
        used=db.one(f'SELECT count(DISTINCT dataset_id) n FROM jobs WHERE dataset_id IN ({placeholders})',ids)['n']
        if used: raise HTTPException(409,f'选中的数据中有 {used} 份正在被任务使用。请先删除相关任务。')
        files=[]
        for row in rows:
            try: name=json.loads(row['stats']).get('uploaded_file','')
            except (TypeError,ValueError): name=''
            if name and Path(name).name==name: files.append(root/'uploads'/name)
            # Validated datasets store only validation statistics. The upload
            # filename is still deterministic: <dataset id>.<extension>.
            files.extend((root/'uploads').glob(row['id']+'.*'))
        files=list(dict.fromkeys(files))
        with db.connect() as conn:
            conn.execute(f'DELETE FROM datasets WHERE id IN ({placeholders})',ids)
        failed=[]
        for path in files:
            try: path.unlink(missing_ok=True)
            except OSError: failed.append(path.name)
        message=f'已删除 {len(ids)} 份数据'
        if failed: message+=f'；{len(failed)} 个原始文件未能删除，可在数据目录中手动清理'
        return {'deleted':len(ids),'message':message}
    @app.get('/api/datasets/{ident}/points')
    def points(ident:str,page:int=Query(1,ge=1),size:int=Query(50,ge=1,le=200),issues:bool=False):
        where="dataset_id=?"+(" AND issues!='[]'" if issues else '')
        return dict(total=db.one(f'SELECT count(*) n FROM points WHERE {where}',(ident,))['n'],rows=db.rows(f'SELECT * FROM points WHERE {where} ORDER BY row_no LIMIT ? OFFSET ?',(ident,size,(page-1)*size)))
    @app.get('/api/sample/{kind}')
    def sample(kind:str):
        if kind not in {'csv','xlsx'}: raise HTTPException(404)
        return FileResponse(resource('resources')/f'test_points.{kind}',filename=f'test_points.{kind}')
    @app.get('/api/templates/{template_name}')
    def template(template_name:str):
        """Download a blank Chinese-column template for either node type."""
        name=template_name.lower()
        ext='xlsx' if name.endswith('.xlsx') else 'csv' if name.endswith('.csv') else ''
        kind=name.rsplit('.',1)[0] if ext else name
        if kind not in {'source-nodes','destinations'} or not ext:
            raise HTTPException(404,'模板不存在')
        # 中文表头便于直接填写；导入器同时接受英文表头和这些中文别名。
        source=['来源ID','客源节点名称','经度','纬度','坐标系','省','市','县区','人口','人口年份','人口统计范围','人口来源','备注']
        destination=['目的地ID','旅游目的地名称','经度','纬度','坐标系','省','市','县区','村名','入口名称','坐标来源','核验状态','备注']
        headers=source if kind=='source-nodes' else destination
        if ext=='csv':
            text=io.StringIO(newline=''); csv.writer(text).writerow(headers); data=text.getvalue().encode('utf-8-sig')
            return StreamingResponse(io.BytesIO(data),media_type='text/csv',headers={'Content-Disposition':f'attachment; filename="{kind}.csv"'})
        from openpyxl import Workbook
        book=Workbook(write_only=True); sheet=book.create_sheet(kind); sheet.append(headers)
        buf=io.BytesIO(); book.save(buf); buf.seek(0)
        return StreamingResponse(buf,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{kind}.xlsx"'})
    @app.get('/api/targets')
    def targets(): return db.rows('SELECT * FROM targets ORDER BY id')

    @app.get('/api/source-nodes')
    def source_nodes(dataset_id:str=''):
        """Return validated source points with optional research population metadata.

        Source nodes remain backed by the existing validated dataset table so
        importing a source file does not create a second copy of the raw file.
        """
        where='p.valid=1'; args=[]
        if dataset_id: where+=' AND p.dataset_id=?'; args.append(dataset_id)
        return db.rows(f'''SELECT p.*,d.name AS dataset_name FROM points p JOIN datasets d ON d.id=p.dataset_id
                           WHERE {where} ORDER BY p.dataset_id,p.id''',args)
    @app.post('/api/targets/bulk-delete')
    def bulk_delete_targets(body:TargetIdsInput):
        ids=list(dict.fromkeys(body.ids))
        with db.connect() as conn:
            conn.execute('BEGIN IMMEDIATE')
            available={row[0] for row in conn.execute('SELECT id FROM targets')}
            if not set(ids).issubset(available):
                raise HTTPException(409,'部分目标已不存在，请刷新列表后重新选择')
            conn.executemany('DELETE FROM targets WHERE id=?',[(ident,) for ident in ids])
        return {'deleted':len(ids)}
    @app.post('/api/targets/convert-crs')
    def convert_targets_crs(body:TargetCrsInput):
        rows=db.rows('SELECT id,lon,lat,crs FROM targets ORDER BY id')
        converted=0
        with db.connect() as conn:
            for row in rows:
                if row['crs']==body.crs: continue
                lon,lat=convert(row['lon'],row['lat'],row['crs'],body.crs)
                conn.execute('UPDATE targets SET lon=?,lat=?,crs=? WHERE id=?',(round(lon,7),round(lat,7),body.crs,row['id']))
                converted+=1
        return {'converted':converted,'unchanged':len(rows)-converted,'crs':body.crs}
    @app.post('/api/targets')
    def add_target(body:TargetInput):
        values=body.model_dump()
        if values['target_id'] and db.one('SELECT id FROM targets WHERE target_id=?',(values['target_id'],)):
            raise HTTPException(409,'稳定目标 ID 已存在，请使用唯一 ID')
        cols=['name','lon','lat','crs','note','target_id','province','city','county','village_name','entrance_name','coordinate_source','verification_status']
        return {'id':db.execute('INSERT INTO targets('+','.join(cols)+') VALUES('+','.join('?' for _ in cols)+')',[values[c] for c in cols])}
    @app.put('/api/targets/{ident}')
    def edit_target(ident:int,body:TargetInput):
        values=body.model_dump(); cols=['name','lon','lat','crs','note','target_id','province','city','county','village_name','entrance_name','coordinate_source','verification_status']
        if values['target_id'] and db.one('SELECT id FROM targets WHERE target_id=? AND id<>?',(values['target_id'],ident)):
            raise HTTPException(409,'稳定目标 ID 已存在，请使用唯一 ID')
        db.execute('UPDATE targets SET '+','.join(f'{c}=?' for c in cols)+' WHERE id=?',[values[c] for c in cols]+[ident]); return {'ok':True}
    @app.delete('/api/targets/{ident}')
    def delete_target(ident:int):
        db.execute('DELETE FROM targets WHERE id=?',(ident,)); return {'ok':True}
    @app.post('/api/datasets/{ident}/as-targets')
    def import_targets(ident:str):
        with db.connect() as conn:
            rows=conn.execute("SELECT * FROM points WHERE dataset_id=? AND valid=1",(ident,)).fetchall()
            existing={str(r[0]) for r in conn.execute('SELECT target_id FROM targets WHERE target_id IS NOT NULL AND target_id!=\'\'')}
            used=set(existing)
            for p in rows:
                original_id=str(p['point_id'] or '').strip()
                target_id=original_id
                id_note=''
                # Keep imports non-destructive while guaranteeing a stable,
                # unique target ID.  Duplicate rows are retained and the
                # deterministic row suffix makes the issue visible in the
                # target list instead of silently merging them.
                if not target_id or target_id in used:
                    target_id=f'{target_id or "ROW"}-row-{p["row_no"]}'
                    while target_id in used: target_id+='-1'
                    id_note='；文件 ID 重复或与已有目标冲突，已按行号生成唯一 ID，请核对'
                used.add(target_id)
                cur=conn.execute('''INSERT INTO targets(target_id,name,lon,lat,crs,note,category,province,city,county,village_name,entrance_name,coordinate_source,verification_status)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',(target_id,p['name'] or original_id,p['lon'],p['lat'],p['crs'],(p['notes'] or '从数据集批量导入')+id_note,p['category'] or '',p['province'],p['city'],p['county'],p['village_name'],p['entrance_name'],p['coordinate_source'] or '文件导入',p['verification_status'] or '待核验'))
                # Keep the complete imported destination row available for
                # task snapshots and export.  The normalized target table is
                # still used for routing; raw JSON is the lossless audit copy.
                conn.execute('INSERT OR REPLACE INTO target_sources(target_id,dataset_id,row_no,raw) VALUES(?,?,?,?)',
                             (cur.lastrowid,ident,p['row_no'],p['raw']))
        return {'ok':True,'added':len(rows)}
    @app.post('/api/jobs/estimate')
    def estimate_job(body:JobInput): return estimate_capacity(db,settings,body)
    @app.post('/api/jobs')
    def create_job(body:JobInput):
        dataset=db.one('SELECT * FROM datasets WHERE id=?',(body.dataset_id,))
        if not dataset: raise HTTPException(400,'请选择数据集')
        source_ids=list(dict.fromkeys(body.source_point_ids or []))
        if source_ids:
            placeholders=','.join('?' for _ in source_ids)
            source_rows=db.rows(f'SELECT id FROM points WHERE dataset_id=? AND valid=1 AND id IN ({placeholders})',[body.dataset_id,*source_ids])
            if len(source_rows)!=len(source_ids):
                raise HTTPException(400,'所选客源地中有不存在或无效的记录，请刷新客源地选择后重试')
            count=len(source_rows)
        else:
            count=db.one('SELECT count(*) n FROM points WHERE dataset_id=? AND valid=1',(body.dataset_id,))['n']
        if not count: raise HTTPException(400,'没有有效点位')
        source_where='dataset_id=? AND valid=1'
        source_args=[body.dataset_id]
        if source_ids:
            source_where+=' AND id IN ('+','.join('?' for _ in source_ids)+')'
            source_args.extend(source_ids)
        source_snapshot=db.rows(f'''SELECT id,point_id,name,lon,lat,crs,wlon,wlat,raw,category,source_id,source_name,province,city,county,
                                         population,population_year,population_scope,population_source,notes,
                                         village_name,entrance_name,coordinate_source,verification_status
                                  FROM points WHERE {source_where} ORDER BY id''',source_args)
        # A legacy validated file may predate stable source IDs.  Construct a
        # deterministic snapshot-only fallback and resolve collisions by the
        # internal row ID; the raw/imported point_id remains untouched.
        used_source_ids=set()
        for source in source_snapshot:
            base=str(source.get('source_id') or source.get('point_id') or source['id'])
            stable=base
            if stable in used_source_ids:
                stable=f'{base}-row-{source["id"]}'
                while stable in used_source_ids:
                    stable+='-1'
            source['source_id']=stable
            used_source_ids.add(stable)
        if body.provider=='amap' and not body.acknowledge_estimate: raise HTTPException(400,'请确认高德当前路线估算不等于指定时刻预测')
        if body.provider=='amap' and not settings.key(): raise HTTPException(400,'请先在设置中保存高德 Web 服务 Key')
        if body.provider=='mock' and body.data_use in {'pilot','formal'}:
            raise HTTPException(400,'模拟服务不能标记为真实试采或正式采集；请改为模拟测试或使用高德真实接口')
        if body.data_use=='formal' and body.collection_mode=='scheduled' and not body.schedule_window_confirmed:
            raise HTTPException(400,'最终方案尚未确认采集窗口；请先确认允许窗口后再创建正式采集任务。真实试采可暂不确认。')
        ids=list(dict.fromkeys(body.target_ids))
        targets=db.rows('SELECT * FROM targets WHERE id IN ('+','.join('?' for _ in ids)+')',ids)
        if len(targets)!=len(ids): raise HTTPException(400,'部分目标地点不存在')
        if targets:
            provenance=db.rows('SELECT target_id,dataset_id,row_no,raw FROM target_sources WHERE target_id IN ('+','.join('?' for _ in ids)+')',ids)
            provenance_by_id={int(row['target_id']):row for row in provenance}
            for target in targets:
                source=provenance_by_id.get(int(target['id']))
                if source:
                    target['source_dataset_id']=source['dataset_id'] or ''
                    target['source_row_no']=source['row_no']
                    target['raw']=source['raw'] or ''
                else:
                    target.setdefault('source_dataset_id','')
                    target.setdefault('source_row_no',None)
                    target.setdefault('raw','')
        for t in targets: t['wlon'],t['wlat']=convert(t['lon'],t['lat'],t['crs'])
        if body.schedule_template==FINAL_TEMPLATE_VERSION:
            # The locked template is a research-scope claim, not merely a
            # seven-row calendar.  Validate the stable-ID universe before a
            # job can carry that version; a subset/pilot is necessarily an
            # adjusted/custom plan.
            source_stable_ids={str(s.get('source_id') or '').strip() for s in source_snapshot}
            target_stable_ids={str(t.get('target_id') or '').strip() for t in targets}
            if body.mode!='exact' or count!=31 or len(targets)!=94:
                raise HTTPException(400,'最终锁定方案必须使用31个客源单元、94个重点村和全部OD；试采或子集请使用自定义/已调整方案')
            expected_source_ids={f'O{i:02d}' for i in range(1,32)}
            expected_target_ids={f'D{i:03d}' for i in range(1,95)}
            if source_stable_ids!=expected_source_ids or target_stable_ids!=expected_target_ids:
                raise HTTPException(400,'最终锁定方案要求使用正式CSV中的完整稳定ID集合（O01—O31、D001—D094）')
        directions=list(dict.fromkeys(body.directions or ['outbound']))
        if not directions:
            raise HTTPException(400,'至少选择去程或返程一个方向')
        if body.collection_mode=='immediate':
            departures=[datetime.now(timezone(timedelta(hours=8))).isoformat(timespec='seconds')]
            departure_metadata={departures[0]:dict(phase='立即采集',day_type=body.day_type,notes='立即采集',directions=directions)}
        else:
            schedule=body.schedule or []
            if schedule:
                # A schedule item owns its date, time list and window.  For
                # compatibility the job-level window remains the default.
                dates=sorted(set(x.date for x in schedule))
                if len(dates)>31: raise HTTPException(400,'定时采集最多设置 31 个日期')
                if any(not x.times for x in schedule): raise HTTPException(400,'每个采集日期至少填写一个时间')
            else:
                dates=sorted(set(body.dates or ([body.date] if body.date else [])))
                if not dates or not body.times: raise HTTPException(400,'定时采集请填写日期和时间')
            departures=[]
            departure_metadata={}
            items=schedule or [dict(date=day,day_type=body.day_type,times=body.times,window_minutes=body.window_minutes,notes='') for day in dates]
            for item in items:
                day=item.date if hasattr(item,'date') else item['date']
                day_type=item.day_type if hasattr(item,'day_type') else item['day_type']
                item_times=item.times if hasattr(item,'times') else item['times']
                item_window=item.window_minutes if hasattr(item,'window_minutes') else item['window_minutes']
                phase=item.phase if hasattr(item,'phase') else item.get('phase','')
                notes=item.notes if hasattr(item,'notes') else item.get('notes','')
                item_directions=item.directions if hasattr(item,'directions') else item.get('directions')
                item_directions=list(dict.fromkeys(item_directions or directions))
                if not item_directions:
                    raise HTTPException(400,'每个采集日期至少选择一个方向')
                for t in sorted(set(item_times)):
                    try: tm=dtime.fromisoformat(t)
                    except (ValueError,TypeError): raise HTTPException(400,'时间格式应为 HH:MM') from None
                    if tm.tzinfo: raise HTTPException(400,'时间使用北京时间 UTC+8，不要附加时区')
                    dep=f'{day.isoformat()}T{tm.strftime("%H:%M")}:00+08:00'
                    departures.append(dep); departure_metadata[dep]=dict(phase=phase,day_type=day_type,window_minutes=item_window,notes=notes,directions=item_directions)
            departures=sorted(set(departures))
            # Check only times belonging to the same date. Different dates can
            # never overlap even when their clock times are equal.
            for day in sorted({d[:10] for d in departures}):
                day_deps=[d for d in departures if d[:10]==day]
                for a,b in zip(day_deps,day_deps[1:]):
                    window=max(int(departure_metadata[a].get('window_minutes',body.window_minutes)),int(departure_metadata[b].get('window_minutes',body.window_minutes)))
                    if (datetime.fromisoformat(b)-datetime.fromisoformat(a)).total_seconds()<window*60:
                        raise HTTPException(400,'同一日期的采集窗口不能重叠，请调整时间或缩短窗口')
            if body.schedule_template==FINAL_TEMPLATE_VERSION:
                if not is_final_plan(schedule):
                    raise HTTPException(400,'最终锁定方案必须为7个指定日期、每天10:00、仅去程；如有改动请使用“已调整方案”而不是最终锁定版本')
                if body.schedule_adjusted:
                    raise HTTPException(400,'最终锁定方案不能同时标记为已调整，请改用自定义方案')
        candidates_per_source=min(body.candidates,len(targets)) if body.mode=='saving' else len(targets)
        total=sum(count*candidates_per_source*len((departure_metadata.get(dep) or {}).get('directions') or directions) for dep in departures)
        if total>5000000: raise HTTPException(400,'单任务上限 500 万组合，请拆分任务')
        if body.collection_mode=='scheduled' and any(datetime.fromisoformat(d).timestamp()<=time.time() for d in departures):
            raise HTTPException(400,'定时采集必须选择未来时间；已过去的时段不能补采')
        estimate=estimate_capacity(db,settings,body)
        if estimate['warnings'] and not body.acknowledge_capacity: raise HTTPException(400,'采集规模存在风险，请查看规模预估并确认后提交')
        schedule_snapshot=[dict(date=x.date.isoformat(),phase=x.phase,day_type=x.day_type,times=list(x.times),window_minutes=x.window_minutes,notes=x.notes,directions=list(x.directions) if x.directions else list(directions)) for x in body.schedule]
        limits=settings.get()
        execution_limits={key:limits.get(key) for key in ('qps','max_concurrency','daily_limit','retries')}
        cfg=dict(capacity_estimate=estimate,execution_limits=execution_limits,day_type=body.day_type,departure_metadata=departure_metadata,directions=directions,study_id=body.study_id,query_order=body.query_order,random_seed=body.random_seed if body.random_seed is not None else secrets.randbelow(2147483648),order_algorithm='sha256-v1',collection_mode=body.collection_mode,window_minutes=body.window_minutes,schedule_window_confirmed=body.schedule_window_confirmed,fresh_collection=True,targets=targets,departures=departures,mode=body.mode,candidates=body.candidates,provider=body.provider,strategy=body.strategy,route_selection_rule='first',point_count=count,source_point_ids=source_ids,api_version='v5/direction/driving' if body.provider=='amap' else 'mock-v1',schedule=schedule_snapshot,schedule_template=body.schedule_template,schedule_adjusted=body.schedule_adjusted,source_snapshot=source_snapshot,data_use=body.data_use)
        jid=uuid.uuid4().hex
        db.execute('INSERT INTO jobs(id,name,created,status,dataset_id,config,total) VALUES(?,?,?,?,?,?,?)',(jid,body.name,now(),'scheduled' if body.collection_mode=='scheduled' else 'queued',body.dataset_id,json.dumps(cfg,ensure_ascii=False),total))
        return get_job(jid)
    @app.get('/api/jobs')
    def jobs():
        rows=db.rows('SELECT id,name,created,status,total,success,failed,message,config FROM jobs ORDER BY created DESC LIMIT 200')
        for row in rows:
            try: row['data_use']=json.loads(row.get('config') or '{}').get('data_use','unmarked') or 'unmarked'
            except (TypeError,ValueError): row['data_use']='unmarked'
            row.pop('config',None)
        return rows
    @app.get('/api/jobs/{jid}')
    def job(jid:str): return get_job(jid)
    @app.post('/api/jobs/{jid}/{action}')
    def control(jid:str,action:str):
        job=get_job(jid)
        if action=='pause': db.execute("UPDATE jobs SET status='paused',message='用户暂停' WHERE id=? AND status IN ('running','queued','scheduled')",(jid,))
        elif action=='resume': db.execute("UPDATE jobs SET status='queued',message='' WHERE id=? AND status IN ('paused','failed')",(jid,))
        elif action=='retry':
            if job['status'] in {'running','queued','scheduled'}: raise HTTPException(409,'请先暂停或等待完成')
            with db.connect() as conn:
                conn.execute("UPDATE results SET status='pending',error=NULL WHERE job_id=? AND status='failed'",(jid,))
                conn.execute("UPDATE jobs SET failed=0,status='queued',message='' WHERE id=?",(jid,))
        else: raise HTTPException(404,'操作不存在')
        return get_job(jid)
    @app.patch('/api/jobs/{jid}')
    def rename(jid:str,body:Rename):
        job=get_job(jid)
        if body.name is None and body.data_use is None: raise HTTPException(400,'至少提供任务名称或数据用途')
        if body.data_use is not None and job['config'].get('provider')=='mock' and body.data_use in {'pilot','formal'}:
            raise HTTPException(400,'模拟服务不能标记为真实试采或正式采集')
        if body.data_use=='formal' and job['config'].get('collection_mode')=='scheduled' and not job['config'].get('schedule_window_confirmed'):
            raise HTTPException(400,'正式采集任务必须先确认允许采集窗口；历史任务配置未记录确认，不能仅靠修改用途补标正式')
        if body.name: db.execute('UPDATE jobs SET name=? WHERE id=?',(body.name,jid))
        if body.data_use is not None:
            cfg=dict(job['config']); old=cfg.get('data_use','unmarked') or 'unmarked'; cfg['data_use']=body.data_use
            db.execute('UPDATE jobs SET config=? WHERE id=?',(json.dumps(cfg,ensure_ascii=False),jid))
            db.execute('INSERT INTO events(created,job_id,level,message) VALUES(?,?,?,?)',(now(),jid,'info',f'数据用途由“{DATA_USE_LABELS.get(old,old)}”改为“{DATA_USE_LABELS.get(body.data_use,body.data_use)}”；仅修改任务元数据，不改变历史观测'))
        return {'ok':True}
    @app.delete('/api/jobs/{jid}')
    def delete_job(jid:str):
        job=get_job(jid)
        if job['status'] in {'queued','running','scheduled'}: raise HTTPException(409,'请先暂停任务')
        db.execute('DELETE FROM jobs WHERE id=?',(jid,)); return {'ok':True}
    @app.post('/api/jobs/bulk-delete')
    def bulk_delete_jobs(body:JobIdsInput):
        ids=list(dict.fromkeys(body.ids))
        placeholders=','.join('?' for _ in ids)
        rows=db.rows(f'SELECT id,status FROM jobs WHERE id IN ({placeholders})',ids)
        if len(rows)!=len(ids): raise HTTPException(404,'部分任务已不存在，请刷新页面后重试')
        active=sum(row['status'] in {'queued','running','scheduled'} for row in rows)
        if active: raise HTTPException(409,f'选中的任务中有 {active} 个正在等待或计算。请先暂停后再删除。')
        with db.connect() as conn:
            conn.execute(f'DELETE FROM jobs WHERE id IN ({placeholders})',ids)
        return {'deleted':len(ids)}
    def query_results(jid,q,status,target,hours,departure,sort,direction,od_direction=''):
        get_job(jid)
        where='job_id=?'; args=[jid]
        if q: where+=' AND (point_id LIKE ? OR point_name LIKE ? OR target_name LIKE ?)'; args.extend(['%'+q+'%']*3)
        if status: where+=' AND status=?'; args.append(status)
        if target: where+=' AND target_name=?'; args.append(target)
        if hours is not None: where+=" AND status='success' AND duration_seconds<=?"; args.append(hours*3600)
        if departure: where+=' AND departure=?'; args.append(departure)
        if od_direction:
            if od_direction not in {'outbound','inbound'}: raise HTTPException(400,'无效出行方向')
            where+=' AND direction=?'; args.append(od_direction)
        if sort not in {'id','straight_km','distance_meters','duration_seconds','point_id','target_name','departure','query_order','requested_at'}: raise HTTPException(400,'无效排序字段')
        if direction not in {'asc','desc'}: raise HTTPException(400,'无效排序方向')
        return where,args,(f'departure,query_order {direction},id' if sort=='query_order' else f'{sort} IS NULL,{sort} {direction},id')
    @app.get('/api/jobs/{jid}/results')
    def results(jid:str,q:str='',status:str='',target:str='',hours:float|None=Query(None,ge=0),departure:str='',sort:str='id',direction:str='asc',od_direction:str='',page:int=Query(1,ge=1),size:int=Query(50,ge=1,le=200)):
        where,args,order=query_results(jid,q,status,target,hours,departure,sort,direction,od_direction)
        return dict(total=db.one(f'SELECT count(*) n FROM results WHERE {where}',args)['n'],rows=db.rows(f'SELECT * FROM results WHERE {where} ORDER BY {order} LIMIT ? OFFSET ?',[*args,size,(page-1)*size]))
    @app.get('/api/jobs/{jid}/observations/{observation_id}/attempts')
    def observation_attempts(jid:str,observation_id:str):
        get_job(jid)
        return db.rows('SELECT * FROM request_attempts WHERE job_id=? AND observation_id=? ORDER BY id',(jid,observation_id))
    @app.get('/api/jobs/{jid}/summary')
    def summary(jid:str,departure:str='',top:int=Query(3,ge=1,le=10),page:int=Query(1,ge=1),size:int=Query(50,ge=1,le=200),od_direction:str=''):
        job=get_job(jid); departure=departure or job['config']['departures'][0]
        return list(summaries(db,jid,departure,top,size,(page-1)*size,od_direction))
    @app.get('/api/jobs/{jid}/research')
    def research(jid:str):
        get_job(jid)
        return research_report(db,jid)
    @app.get('/api/jobs/{jid}/export')
    def export(jid:str,kind:str='details',q:str='',status:str='',target:str='',hours:float|None=Query(None,ge=0),departure:str='',sort:str='id',direction:str='asc',od_direction:str='',include_raw:bool=True):
        job=get_job(jid)
        # Research, paper-analysis and quality exports are read-only snapshots
        # and explicitly label an unfinished task as a stage export.  The
        # detail/summary/package exports retain the older consistency guard.
        if kind not in {'research','paper','quality'} and job['status'] in {'queued','running','scheduled'}: raise HTTPException(409,'请暂停任务或等待完成后导出一致结果')
        path=root/'exports'/f'{jid}-{uuid.uuid4().hex}.xlsx'
        if kind=='research': export_research(research_report(db,jid),path)
        elif kind=='package':
            path=root/'exports'/f'{jid}-{uuid.uuid4().hex}.zip'; export_package(db,path,jid,include_raw=include_raw)
        elif kind=='summary': export_summary(db,path,jid,departure or job['config']['departures'][0],od_direction)
        elif kind=='paper': export_paper(db,path,jid)
        elif kind=='quality': export_quality(db,path,jid)
        elif kind=='details':
            where,args,order=query_results(jid,q,status,target,hours,departure,sort,direction,od_direction)
            export_details(db,path,where,args,order)
        else: raise HTTPException(400,'未知导出类型')
        stamp=datetime.now(timezone(timedelta(hours=8))).strftime('%Y%m%d-%H%M%S'); short=jid[:8]; safe_name=re.sub(r'[\\/:*?"<>|]+','_',job['name']).strip()[:60] or '任务'
        names={'paper':f'论文分析数据_{safe_name}_{stamp}_{short}.xlsx','quality':f'采集质量报告_{safe_name}_{stamp}_{short}.xlsx','package':f'完整采集审计包_{safe_name}_{stamp}_{short}.zip','summary':f'起点汇总_{safe_name}_{stamp}_{short}.xlsx','details':f'路线明细_{safe_name}_{stamp}_{short}.xlsx','research':f'采集完整性报告_{safe_name}_{stamp}_{short}.xlsx'}
        return FileResponse(path,filename=names.get(kind,f'交通可达性_{kind}_{short}.xlsx'))
    @app.post('/api/shutdown')
    async def shutdown():
        worker.stop_event.set()
        if on_shutdown: asyncio.get_running_loop().call_later(.3,on_shutdown)
        return {'ok':True}
    return app
