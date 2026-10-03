import csv
import io
import json
import math
from pathlib import Path
from openpyxl import load_workbook
from utils.geo import convert

ALIASES={
    # The first aliases are deliberately the canonical research columns.  A
    # table can contain both WGS84 and GCJ-02 coordinates; choosing aliases by
    # priority (rather than by their position in the file) prevents the
    # routing coordinate from silently being replaced by the audit coordinate.
    'point_id':['origin_id','destination_id','source_id','id','客源地编号','重点村编号','客源节点编号','目的地编号','编号','序号','点位id','来源id','客源点id','客源节点id','目标id','目的地id'],
    'name':['origin_name','destination_name','source_name','name','客源地名称','重点村名称','村名','选定POI名称','客源市场单元','客源点名称','客源节点名称','目标名称','目的地名称','旅游目的地名称','名称','点位名称','地点','来源名称'],
    'lon':['gcj02_lon','final_route_lon','origin_lon_gcj02','final_lon_gcj02','客源地经度（GCJ-02）','最终路线经度（GCJ-02）','POI经度（GCJ-02）','导航入口经度（GCJ-02）','经度（GCJ-02）','longitude','lng','lon','经度','wgs84_lon'],
    'lat':['gcj02_lat','final_route_lat','origin_lat_gcj02','final_lat_gcj02','客源地纬度（GCJ-02）','最终路线纬度（GCJ-02）','POI纬度（GCJ-02）','导航入口纬度（GCJ-02）','纬度（GCJ-02）','latitude','lat','纬度','wgs84_lat'],
    # A file may declare one CRS for the whole table (the wizard's ``crs``
    # value) or provide a coordinate_system/坐标系 column per row.  The latter
    # is useful when combining carefully documented sources; blank cells fall
    # back to the file-level choice.
    'coordinate_system':['routing_coord_system','coordinate_system','路线查询坐标系','人口重心坐标系','坐标系','crs','坐标参考系','coordinate crs'],
    'province':['province','省','省份','省级行政区'], 'city':['market_unit','city','市','城市','客源市场单元','市级行政区'], 'county':['county','县','区','县区','县级行政区'],
    'population':['urban_population','urban_population_ghsl','population','城市人口（GHSL）','人口','常住人口','城镇人口'], 'population_year':['population_year','ghs_pop_year','人口年份','GHS人口年份','统计年份'],
    'population_scope':['population_scope','population_weight_type','范围类型','人口口径','统计范围'], 'population_source':['population_source','人口来源'],
    'notes':['review_note','notes','note','备注','说明'],
    'village_name':['village_name','村名','重点村名称','乡村旅游重点村'],
    'entrance_name':['route_poi_name','selected_poi_name','entrance_name','选定POI名称','入口名称','游客中心','景区入口'],
    'coordinate_source':['coordinate_source','坐标来源','坐标确定来源','位置来源'],
    'verification_status':['verification_status','核验状态','验证状态','是否可用于路线查询','ready_for_routing'],
    # These are retained in the raw row in all cases.  Mapping them to the
    # existing metadata columns also makes the quality information visible in
    # the data-management preview without reintroducing institution
    # classification or deduplication.
    'category':['selection_level','地点等级','质量等级','可信度等级'],
    'quality_flag':['quality_flag','质量标记','质量检查状态'],
    'route_poi_type':['route_poi_type','selected_poi_type','选定POI类型','导航地点类型','入口类型'],
    'official_batch':['official_batch','官方公布批次','公布批次','批次'],
    'coordinate_method':['coordinate_method','gcj_conversion_method','GCJ-02转换方法','坐标方法','坐标确定方法'],
    'boundary_definition':['boundary_definition','边界定义','纳入范围说明'],
    'included_area':['included_area','纳入区域','纳入范围'],
    'ready_for_routing':['ready_for_routing','是否可用于路线查询','可路由','可用于路由']
}

def canonical_crs(value):
    key=str(value or '').strip().upper().replace('_','').replace('-','').replace(' ','')
    return {'WGS84':'WGS84','GCJ02':'GCJ-02','BD09':'BD-09'}.get(key)

def read_rows(path):
    if Path(path).suffix.lower()=='.xlsx':
        book=load_workbook(path,read_only=True,data_only=True)
        try:
            yield from book.active.iter_rows(values_only=True)
        finally:
            book.close()
    else:
        raw=Path(path).read_bytes()
        decoded=None
        for encoding in ('utf-8-sig','gb18030','utf-16'):
            try:
                decoded=raw.decode(encoding)
                break
            except UnicodeError:
                pass
        if decoded is None:
            raise ValueError('无法识别 CSV 编码，请另存为 UTF-8 CSV')
        try:
            dialect=csv.Sniffer().sniff(decoded[:8192],delimiters=',;\t')
        except csv.Error:
            dialect=csv.excel
        yield from csv.reader(io.StringIO(decoded),dialect)

def headers(path):
    rows=read_rows(path)
    try:
        first=next(rows,None)
    finally:
        rows.close()
    if not first:
        raise ValueError('文件为空')
    cols=[str(v).strip() if v is not None else f'未命名列{i+1}' for i,v in enumerate(first)]
    if len(set(cols))!=len(cols):
        raise ValueError('表头存在重复列名，请先重命名')
    normalized=[v.lower().replace(' ','') for v in cols]
    # Iterate aliases first so a canonical column (e.g. gcj02_lon) wins even
    # when an audit column (e.g. wgs84_lon) appears earlier in the CSV.
    mapping={k:next((normalized.index(alias.lower().replace(' ','')) for alias in aliases if alias.lower().replace(' ','') in normalized),None) for k,aliases in ALIASES.items()}
    return cols,mapping

def ingest(db,dataset,path,mapping,crs):
    cols,_=headers(path)
    if mapping.get('lon') is None or mapping.get('lat') is None:
        raise ValueError('请选择经度和纬度字段')
    if mapping['lon']==mapping['lat']:
        raise ValueError('经度和纬度不能选择同一列')
    for value in mapping.values():
        if value is not None and (not isinstance(value,int) or not 0<=value<len(cols)):
            raise ValueError('字段映射无效')
    stats=dict(total=0,valid=0,invalid=0,duplicate=0,empty=0)
    seen_ids=set(); used_source_ids=set(); seen_coords=set()
    with db.connect() as conn:
        conn.execute('DELETE FROM points WHERE dataset_id=?',(dataset,))
        rows=read_rows(path)
        next(rows)
        for number,row in enumerate(rows,2):
            stats['total']+=1
            values=list(row)
            def field(key):
                i=mapping.get(key)
                return '' if i is None or i>=len(values) or values[i] is None else str(values[i]).strip()
            pid=field('point_id') or f'ROW-{number}'
            name=field('name')
            issues=[]; lon=lat=wlon=wlat=None; good=True; row_crs=crs
            declared_crs=field('coordinate_system')
            if declared_crs:
                row_crs=canonical_crs(declared_crs)
                if row_crs is None:
                    issues.append('坐标系无效（应为 WGS84、GCJ-02 或 BD-09）'); good=False
            if not any(v is not None and str(v).strip() for v in values):
                issues.append('空行'); stats['empty']+=1; good=False
            ready_value=field('ready_for_routing').strip().lower()
            if ready_value and ready_value not in {'true','1','yes','y','是','可路由','ok'}:
                issues.append('文件标记为不可路由'); good=False
            try:
                if not field('lon') or not field('lat'):
                    raise ValueError('经纬度缺失')
                lon,lat=float(field('lon')),float(field('lat'))
                if not math.isfinite(lon) or not math.isfinite(lat):
                    raise ValueError('经纬度不是有限数字')
                if abs(lon)<=90 and 90<abs(lat)<=180:
                    issues.append('经纬度疑似写反')
                if abs(lon)>180: issues.append('经度超范围'); good=False
                if abs(lat)>90: issues.append('纬度超范围'); good=False
                if good: wlon,wlat=convert(lon,lat,row_crs)
            except (ValueError,TypeError) as exc:
                good=False; issues.append(str(exc) if '经纬度' in str(exc) else '经纬度非数字')
                if lon is not None and not math.isfinite(lon): lon=None
                if lat is not None and not math.isfinite(lat): lat=None
            duplicate=False
            if pid in seen_ids:
                issues.append('重复 ID'); duplicate=True
            seen_ids.add(pid)
            # Keep the user's point_id untouched for audit, but give every
            # valid row a unique, deterministic source_id for OD joins.  The
            # suffix is based on the original file row, so re-importing the
            # same file reproduces the identifier exactly.
            stable_source_id=pid
            if stable_source_id in used_source_ids:
                # A user-supplied ID can itself collide with a deterministic
                # suffix generated for an earlier duplicate. Keep every
                # source_id unique without changing the audited point_id.
                stable_source_id=f'{pid}-row-{number}'
                while stable_source_id in used_source_ids:
                    stable_source_id+='-1'
                if not duplicate:
                    issues.append('稳定客源 ID 冲突，已按行号生成唯一 ID')
                    duplicate=True
            used_source_ids.add(stable_source_id)
            if lon is not None and lat is not None:
                key=(lon,lat)
                if key in seen_coords: issues.append('重复坐标'); duplicate=True
                seen_coords.add(key)
            stats['duplicate']+=int(duplicate)
            stats['valid' if good else 'invalid']+=1
            population=None
            if field('population'):
                try:
                    population=float(field('population'))
                    if not math.isfinite(population) or population<0: raise ValueError
                except (ValueError,TypeError):
                    # Population is research metadata, not a route coordinate;
                    # retain the row as valid but preserve the problem in issues.
                    population=None; issues.append('人口不是有效非负数字')
            population_year=None
            if field('population_year'):
                try: population_year=int(float(field('population_year')))
                except (ValueError,TypeError): population_year=None; issues.append('人口年份不是整数')
            # Keep the research quality/traceability fields visible in the
            # existing metadata columns while retaining the complete original
            # row in ``raw``.  This is metadata only; no institution grouping
            # or automatic deduplication is performed.
            note=field('notes')
            quality_parts=[]
            for key,label in [('quality_flag','质量标记'),('route_poi_type','导航地点类型'),('official_batch','公布批次'),('coordinate_method','坐标方法')]:
                value=field(key)
                if value: quality_parts.append(f'{label}：{value}')
            if quality_parts:
                note='；'.join([note,*quality_parts]) if note else '；'.join(quality_parts)
            verification=field('verification_status') or ready_value
            conn.execute('''INSERT INTO points(dataset_id,row_no,raw,point_id,name,lon,lat,crs,wlon,wlat,valid,issues,category,source_id,source_name,province,city,county,population,population_year,population_scope,population_source,notes,village_name,entrance_name,coordinate_source,verification_status)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',(
                dataset,number,json.dumps(values,ensure_ascii=False,default=str),pid,name,lon,lat,row_crs,wlon,wlat,int(good),json.dumps(issues,ensure_ascii=False),
                field('category'),stable_source_id,name,field('province'),field('city'),field('county'),population,population_year,field('population_scope'),field('population_source'),note,field('village_name'),field('entrance_name'),field('coordinate_source'),verification))
        conn.execute('UPDATE datasets SET crs=?,stats=? WHERE id=?',(crs,json.dumps(stats),dataset))
    return stats
