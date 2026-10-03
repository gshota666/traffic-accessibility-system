import io
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import create_app


def _app(tmp_path):
    return create_app(tmp_path / 'paper-export-data', token='paper-token', run_worker=False)


def _dataset(client):
    body = ('source_id,source_name,longitude,latitude,province,population,population_year,population_scope\n'
            'O01,来源甲,114.30,30.50,湖北,1000000,2020,城区\n').encode()
    uploaded = client.post('/api/imports', files={'file': ('sources.csv', body)}).json()
    checked = client.post(f"/api/imports/{uploaded['id']}/validate",
                          json={'mapping': uploaded['mapping'], 'crs': 'WGS84'})
    assert checked.status_code == 200
    return uploaded['id']


def _finished_job(app, client, data_use='pilot'):
    dataset = _dataset(client)
    created = client.post('/api/jobs', json={
        'name': '真实试采导出测试', 'dataset_id': dataset, 'target_ids': [1, 2],
        'provider': 'mock', 'directions': ['outbound', 'inbound'],
        'data_use': 'unmarked', 'collection_mode': 'immediate',
    })
    assert created.status_code == 200, created.text
    job = created.json()
    app.state.db.execute("UPDATE jobs SET status='running' WHERE id=?", (job['id'],))
    app.state.worker.run(app.state.db.one('SELECT * FROM jobs WHERE id=?', (job['id'],)))
    # This is an offline fixture transformation, not an API call.  It lets
    # the export classifier exercise the real-Amap/pilot branch safely.
    app.state.db.execute("UPDATE results SET provider='amap',traffic_type='current_estimate',cache_hit=0 WHERE job_id=?", (job['id'],))
    import json
    snapshot = app.state.db.one('SELECT config FROM jobs WHERE id=?', (job['id'],))
    config = json.loads(snapshot['config']); config['provider'] = 'amap'; config['api_version'] = 'v5/direction/driving'
    app.state.db.execute('UPDATE jobs SET config=? WHERE id=?', (json.dumps(config, ensure_ascii=False), job['id']))
    app.state.db.execute("UPDATE jobs SET status='completed' WHERE id=?", (job['id'],))
    assert client.patch(f"/api/jobs/{job['id']}", json={'data_use': data_use}).status_code == 200
    return job['id']


def test_paper_workbook_has_three_chinese_sheets_and_keeps_route_ids(tmp_path):
    app = _app(tmp_path)
    with TestClient(app, headers={'x-session-token': 'paper-token'}) as client:
        jid = _finished_job(app, client)
        response = client.get(f'/api/jobs/{jid}/export?kind=paper')
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.content), data_only=False)
        assert workbook.sheetnames == ['论文分析主表', '数据质量说明', '字段说明']
        main = workbook['论文分析主表']
        assert main.max_row == 5  # 1 source × 2 targets × 2 directions + header
        headers = [cell.value for cell in main[1]]
        assert headers == ['观测编号','客源地编号','客源地名称','重点村编号','重点村名称','计划日期',
                           '研究情景','方向','城市人口（人）','预计驾车时间（分钟）','驾车距离（千米）','直线距离（千米）']
        removed = {'数据用途','研究阶段','日期类型','相对计划时间偏移（秒）','是否跨窗口响应','数据类型',
                   '人口加权指标资格','人口来源','人口单位','质量检查结果','正式分析纳入状态','说明',
                   '方案版本','是否调整方案','目的地坐标来源','目的地核验状态','正式计划资格',
                   '任务编号','采集批次编号','计划采集时间','周末组别','客源省份','人口年份','人口口径',
                   '目的地省份','城市','县区','乡镇','坐标可信度等级','实际请求时间（北京时间）','响应时间（北京时间）',
                   '地图服务','路线策略','路线选择规则','客源地经度','客源地纬度','客源地坐标系',
                   '重点村经度','重点村纬度','重点村坐标系'}
        assert not removed.intersection(headers)
        assert main.cell(2, headers.index('方向') + 1).value in {'去程', '返程'}
        assert isinstance(main.cell(2, headers.index('预计驾车时间（分钟）') + 1).value, float)
        assert isinstance(main.cell(2, headers.index('驾车距离（千米）') + 1).value, float)
        assert isinstance(main.cell(2, headers.index('直线距离（千米）') + 1).value, float)
        quality_text = ' '.join(str(cell.value or '') for row in workbook['数据质量说明'].iter_rows(min_row=1, max_col=3) for cell in row)
        assert '试采验证，不纳入正式分析' in quality_text
        for label in ('数据用途','研究阶段','相对计划时间偏移（秒）','质量检查结果','方案版本','正式计划资格'):
            assert label in quality_text
        quality_headers = [cell.value for row in workbook['数据质量说明'].iter_rows() for cell in row if cell.value]
        assert '人口来源' in quality_headers and '目的地核验状态' in quality_headers


def test_audit_package_has_chinese_files_and_raw_alias_without_rewriting_json(tmp_path):
    app = _app(tmp_path)
    with TestClient(app, headers={'x-session-token': 'paper-token'}) as client:
        jid = _finished_job(app, client)
        response = client.get(f'/api/jobs/{jid}/export?kind=package&include_raw=false')
        assert response.status_code == 200
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            names = set(archive.namelist())
            assert {'采集计划.csv', '路线观测明细.csv', '数据质量报告.csv', '字段说明.csv', '使用说明.txt'} <= names
            assert '原始响应索引.csv' in names
            assert 'raw_responses/' not in ''.join(names)
            assert '计划采集时间' in archive.read('采集计划.csv').decode('utf-8-sig')


def test_mock_cannot_be_marked_as_real_purpose(tmp_path):
    app = _app(tmp_path)
    with TestClient(app, headers={'x-session-token': 'paper-token'}) as client:
        dataset = _dataset(client)
        response = client.post('/api/jobs', json={
            'name': '模拟任务', 'dataset_id': dataset, 'target_ids': [1],
            'provider': 'mock', 'data_use': 'formal',
        })
        assert response.status_code == 400


def test_unfinished_task_can_export_a_labeled_stage_snapshot(tmp_path):
    app = _app(tmp_path)
    with TestClient(app, headers={'x-session-token': 'paper-token'}) as client:
        dataset = _dataset(client)
        created = client.post('/api/jobs', json={
            'name': '阶段性导出', 'dataset_id': dataset, 'target_ids': [1],
            'provider': 'mock', 'data_use': 'unmarked',
        })
        assert created.status_code == 200
        jid = created.json()['id']
        app.state.db.execute("UPDATE jobs SET status='running' WHERE id=?", (jid,))
        response = client.get(f'/api/jobs/{jid}/export?kind=quality')
        assert response.status_code == 200
