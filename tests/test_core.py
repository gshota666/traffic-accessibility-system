import io
import json
import time
from pathlib import Path
import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook,Workbook
from backend.app import create_app
from providers.amap import AmapProvider
from providers.mock import MockProvider
from providers.base import RouteError
from services.worker import cache_key,Worker
from utils.geo import haversine,convert,valid
from utils.paths import resource,data_dir

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr('services.settings.keyring.get_password',lambda *args: None)
    app=create_app(tmp_path/'中文 数据目录',token='test-token',run_worker=False)
    with TestClient(app,headers={'x-session-token':'test-token'}) as c:
        yield c

def imported(c,kind='csv',content=None):
    data=content if content is not None else (resource('resources')/f'test_points.{kind}').read_bytes()
    r=c.post('/api/imports',files={'file':('中文 测试文件.'+kind,data)})
    assert r.status_code==200,r.text
    uploaded=r.json()
    r=c.post(f'/api/imports/{uploaded["id"]}/validate',json={'mapping':uploaded['mapping'],'crs':'GCJ-02'})
    assert r.status_code==200,r.text
    return uploaded['id'],r.json()

def job(c,dataset,**extra):
    body=dict(name='测试分析',dataset_id=dataset,target_ids=[1,2,3],date='2026-09-15',times=['08:00','12:00'],provider='mock')
    body.update(extra)
    r=c.post('/api/jobs',json=body)
    assert r.status_code==200,r.text
    return r.json()

def run(c,j):
    c.app.state.db.execute("UPDATE jobs SET status='running' WHERE id=?",(j['id'],))
    c.app.state.worker.run(c.app.state.db.one('SELECT * FROM jobs WHERE id=?',(j['id'],)))
    return c.get('/api/jobs/'+j['id']).json()

@pytest.mark.parametrize('kind',['csv','xlsx'])
def test_import_chinese(client,kind):
    ident,s=imported(client,kind)
    assert s==dict(total=17,valid=13,invalid=4,duplicate=1,empty=0)
    r=client.get(f'/api/datasets/{ident}/points?issues=true').json()
    assert r['total']==5
    assert any('疑似写反' in p['issues'] for p in r['rows'])
    assert client.app.state.db.path.exists()

@pytest.mark.parametrize('encoding',['utf-8','utf-8-sig','gb18030'])
def test_csv_encodings(client,encoding):
    _,s=imported(client,content='ID,名称,lng,lat\n001,武汉,114.3,30.5\n'.encode(encoding))
    assert s['valid']==1

def test_validation(client):
    ident,s=imported(client,content='ID,名称,lon,lat\na,缺失,,31\nb,非数字,x,y\nc,无穷,NaN,20\nd,范围,181,0\ne,正常,0,0\nf,正常,0,0\n,,,\n'.encode())
    assert s['valid']==2 and s['invalid']==5 and s['duplicate']==1 and s['empty']==1
    assert len(client.get(f'/api/datasets/{ident}/points').json()['rows'])==7

def test_bad_file_and_mapping(client):
    assert client.post('/api/imports',files={'file':('bad.xlsx',b'broken')}).status_code==400
    assert client.post('/api/imports',files={'file':('bad.csv',b'ID,ID\n1,2')}).status_code==400
    r=client.post('/api/imports',files={'file':('test.csv',b'x,y\n1,2')}).json()
    assert client.post(f'/api/imports/{r["id"]}/validate',json={'mapping':{'lon':0,'lat':0},'crs':'WGS84'}).status_code==400

def test_geo():
    assert haversine((0,0),(0,0))==0
    assert haversine((0,0),(1,0))==pytest.approx(111.195,abs=.01)
    assert haversine((0,0),(180,0))==pytest.approx(20015.114,abs=.01)
    for crs in ['GCJ-02','BD-09']:
        g=convert(116.397,39.908,'WGS84',crs)
        assert convert(*g,crs)==pytest.approx((116.397,39.908),abs=2e-6)
    assert convert(-73,40,'WGS84','GCJ-02')==(-73,40)
    assert not valid(float('inf'),0)

def test_provider_mock_and_amap():
    mock=MockProvider().route((0,0),(1,0),'2026-01-01','32')
    assert mock.traffic_type=='mock' and mock.distance_meters>111000
    def handler(request):
        assert request.url.params['show_fields']=='cost'
        assert 'departure' not in request.url.params
        return httpx.Response(200,json={'status':'1','infocode':'10000','route':{'paths':[{'distance':'1234','cost':{'duration':'123'}}]}})
    with httpx.Client(transport=httpx.MockTransport(handler)) as c:
        r=AmapProvider('secret',c).route((114.3,30.5),(115,31),'2026-09-15T08:00','32')
        assert r.distance_meters==1234 and r.duration_seconds==123 and r.traffic_type=='current_estimate'

@pytest.mark.parametrize('response,retry,fatal',[(httpx.Response(429),True,False),(httpx.Response(503),True,False),(httpx.Response(200,json={'status':'0','infocode':'10001'}),False,True),(httpx.Response(200,json={'status':'1','route':{'paths':[]}}),False,False)])
def test_provider_errors(response,retry,fatal):
    with httpx.Client(transport=httpx.MockTransport(lambda _:response)) as c:
        with pytest.raises(RouteError) as err: AmapProvider('secret',c).route((0,0),(1,1),'','32')
        assert err.value.retryable==retry and err.value.fatal==fatal
        assert 'secret' not in str(err.value)

def test_cache_key():
    args=[(1,2),(3,4),'WGS84','2026-09-15T08:00','32','mock']
    key=cache_key(*args)
    for i,new in enumerate([(2,2),(3,5),'GCJ-02','2026-09-15T12:00','33','amap']):
        changed=args.copy(); changed[i]=new
        assert cache_key(*changed)!=key

def test_task_cache_pagination_export_summary(client):
    ident,s=imported(client)
    j=run(client,job(client,ident))
    assert j['status']=='completed' and j['success']==39
    result=client.get(f'/api/jobs/{j["id"]}/results?size=10&page=2&sort=duration_seconds').json()
    assert len(result['rows'])==10 and result['total']==39
    assert all(r['request_coords'] and r['traffic_type']=='mock' for r in result['rows'])
    j2=run(client,job(client,ident))
    assert client.app.state.db.one('SELECT count(*) n FROM results WHERE job_id=? AND cache_hit=1',(j2['id'],))['n']==0
    assert client.get('/api/usage').json()==[]
    r=client.get(f'/api/jobs/{j["id"]}/export')
    assert r.status_code==200
    wb=load_workbook(io.BytesIO(r.content)); ws=wb.active
    assert ws.max_row==40 and ws['A2'].value=='A001'
    assert ws['Q2'].value=='模拟测试数据'
    assert ws['O2'].is_date
    summary_export=client.get(f'/api/jobs/{j["id"]}/export?kind=summary')
    assert summary_export.status_code==200
    summary_wb=load_workbook(io.BytesIO(summary_export.content), read_only=True)
    summary_ws=summary_wb['起点汇总']
    summary_headers=[cell.value for cell in next(summary_ws.iter_rows(max_row=1))]
    assert '1小时内成功重点村数' in summary_headers
    assert '2小时内成功重点村数' in summary_headers
    assert '3小时内成功重点村数' in summary_headers
    assert '4小时内成功重点村数' in summary_headers
    summary=client.get(f'/api/jobs/{j["id"]}/summary').json()
    assert len(summary)==13 and len(summary[0]['top'])==3
    assert set(summary[0]['reachable'])=={'1','2','3','4'}
    assert client.get(f'/api/jobs/{j["id"]}/results?sort=DROP%20TABLE').status_code==400
    assert client.post(f'/api/imports/{ident}/validate',json={'mapping':{'lon':2,'lat':3},'crs':'WGS84'}).status_code==409

def test_saving_mode(client):
    ident,_=imported(client); j=run(client,job(client,ident,mode='saving',candidates=1))
    assert j['total']==13 and j['success']==13
    assert client.get(f'/api/jobs/{j["id"]}/summary').json()[0]['scope'].startswith('候选入口预筛选结果')

def test_resume_does_not_recompute_success(client):
    ident,_=imported(client); j=job(client,ident)
    db=client.app.state.db;worker=client.app.state.worker
    db.execute("UPDATE jobs SET status='running' WHERE id=?",(j['id'],))
    worker.prepare(j,j['config'])
    row=db.one('SELECT id FROM results WHERE job_id=? ORDER BY id LIMIT 1',(j['id'],))
    db.execute("UPDATE results SET status='success',distance_meters=42,duration_seconds=43 WHERE id=?",(row['id'],))
    db.execute('UPDATE jobs SET success=1 WHERE id=?',(j['id'],))
    worker.start()
    assert db.one('SELECT status FROM jobs WHERE id=?',(j['id'],))['status']=='paused'
    assert client.post(f'/api/jobs/{j["id"]}/resume').status_code==200
    until=time.monotonic()+120
    while time.monotonic()<until:
        result=client.get('/api/jobs/'+j['id']).json()
        if result['status']=='completed': break
        time.sleep(.05)
    worker.close()
    assert result['success']==39
    assert db.one('SELECT distance_meters FROM results WHERE id=?',(row['id'],))['distance_meters']==42
    assert db.one('SELECT count(*) n FROM results WHERE job_id=?',(j['id'],))['n']==39

def test_single_failure_and_retry(client,monkeypatch):
    ident,_=imported(client); j=job(client,ident,target_ids=[1],times=['08:00'])
    original=MockProvider.route
    def fail(self,*args): raise RouteError('无路线')
    monkeypatch.setattr(MockProvider,'route',fail)
    j=run(client,j); assert j['status']=='partial' and j['failed']==13
    monkeypatch.setattr(MockProvider,'route',original)
    client.post(f'/api/jobs/{j["id"]}/retry'); j=run(client,j)
    assert j['success']==13 and j['failed']==0

def test_formula_injection(client):
    ident,_=imported(client,content='ID,名称,lon,lat\n=1+1,测试,114,30\n'.encode())
    j=run(client,job(client,ident,target_ids=[1],times=['08:00']))
    r=client.get(f'/api/jobs/{j["id"]}/export'); cell=load_workbook(io.BytesIO(r.content)).active['A2']
    assert cell.value=='=1+1' and cell.data_type=='s'

def test_local_security(client):
    assert client.get('/api/jobs',headers={'x-session-token':'bad'}).status_code==401
    assert client.get('/api/jobs',headers={'host':'evil.example'}).status_code==403
    assert client.post('/api/shutdown',headers={'origin':'https://evil.example'}).status_code==403
    assert 'api_key' not in client.get('/api/settings').json()

def test_target_snapshot(client):
    ident,_=imported(client); j=job(client,ident)
    client.put('/api/targets/1',json=dict(name='新北京',lon=110,lat=30,crs='WGS84',note=''))
    client.delete('/api/targets/2')
    j=run(client,j)
    assert j['config']['targets'][0]['name']=='北京' and j['success']==39

def test_rate_budget_pause_and_fatal(client,monkeypatch):
    from services.settings import Settings
    from providers.base import RouteResult
    monkeypatch.setattr(Settings,'key',lambda _: 'test-key')
    cfg=dict(provider='amap',qps=20,daily_limit=2,retries=0,strategy='32')
    client.app.state.worker.settings.save(cfg)
    calls=[]
    def route(self,origin,dest,*_):
        calls.append(time.monotonic())
        assert all(v==round(v,6) for v in (*origin,*dest))
        return RouteResult(123,45,'current_estimate','amap','10000')
    monkeypatch.setattr(AmapProvider,'route',route)
    ident,_=imported(client);j=run(client,job(client,ident,provider='amap',acknowledge_estimate=True,acknowledge_capacity=True))
    assert j['status']=='paused' and j['success']==2 and len(calls)==2
    assert calls[1]-calls[0]>=.045
    assert client.get('/api/usage').json()[0]['calls']==2
    cfg['daily_limit']=10;client.app.state.worker.settings.save(cfg)
    def invalid(*_):raise RouteError('Key 无效',fatal=True)
    monkeypatch.setattr(AmapProvider,'route',invalid)
    j=run(client,j)
    assert j['status']=='paused' and j['success']==2 and j['failed']==0

def test_retry_transient_then_success(client,monkeypatch):
    attempts=[]
    original=MockProvider.route
    def flaky(self,*args):
        attempts.append(1)
        if len(attempts)==1:raise RouteError('网络超时',retryable=True)
        return original(self,*args)
    monkeypatch.setattr(MockProvider,'route',flaky)
    ident,_=imported(client,content='ID,名称,lon,lat\na,测试,114,30\n'.encode())
    j=run(client,job(client,ident,target_ids=[1],times=['08:00']))
    assert j['success']==1 and len(attempts)==2

def test_cache_expiry_and_provenance(client,monkeypatch):
    ident,_=imported(client,content='ID,名称,lon,lat\na,测试,114,30\n'.encode())
    first=run(client,job(client,ident,target_ids=[1],times=['08:00']))
    old=client.get(f'/api/jobs/{first["id"]}/results').json()['rows'][0]
    assert old['fetched_at']
    original=MockProvider.route;calls=[]
    def tracked(self,*args):calls.append(1);return original(self,*args)
    monkeypatch.setattr(MockProvider,'route',tracked)
    second=run(client,job(client,ident,target_ids=[1],times=['08:00']))
    cached=client.get(f'/api/jobs/{second["id"]}/results').json()['rows'][0]
    assert len(calls)==1 and cached['cache_hit']==0 and cached['requested_at']
    client.app.state.db.execute('UPDATE cache SET expires=0')
    run(client,job(client,ident,target_ids=[1],times=['08:00']))
    assert len(calls)==2

def test_target_coordinate_error_is_readable_and_does_not_write(client):
    count=len(client.get('/api/targets').json())
    response=client.post('/api/targets',json={'name':'武汉轻工大学','lon':114240124,'lat':30.636285,'crs':'GCJ-02'})
    assert response.status_code==422
    assert response.json()['detail']=='经度必须是 -180 到 180 之间的数字，请检查小数点。'
    assert len(client.get('/api/targets').json())==count
    response=client.post('/api/targets',json={'name':'坐标输入测试','lon':114.240124,'lat':30.636285,'crs':'GCJ-02'})
    assert response.status_code==200
    saved=client.get('/api/targets').json()[-1]
    assert saved['lon']==114.240124 and saved['lat']==30.636285


def test_coordinate_system_guidance_is_present():
    page=(resource('frontend')/'app.js').read_text(encoding='utf-8')
    assert '请选择上传文件原本采用的坐标系' in page
    assert '不是希望转换成的坐标系' in page
    assert 'GPS 设备、OpenStreetMap' in page
    assert '高德地图、腾讯地图' in page
    assert '百度地图使用的坐标系' in page
    assert '软件会保留原始坐标' in page

def test_convert_all_target_coordinate_systems_preserves_locations(client):
    original=client.get('/api/targets').json()
    expected={t['id']:convert(t['lon'],t['lat'],t['crs'],'WGS84') for t in original}
    response=client.post('/api/targets/convert-crs',json={'crs':'WGS84'})
    assert response.status_code==200
    assert response.json()['converted']==31
    converted=client.get('/api/targets').json()
    assert all(t['crs']=='WGS84' for t in converted)
    for target in converted:
        assert (target['lon'],target['lat'])==pytest.approx(expected[target['id']],abs=2e-6)
    response=client.post('/api/targets/convert-crs',json={'crs':'WGS84'}).json()
    assert response=={'converted':0,'unchanged':31,'crs':'WGS84'}
    assert client.post('/api/targets/convert-crs',json={'crs':'BAD'}).status_code==422
    assert client.post('/api/targets/convert-crs',json={'crs':'BD-09'}).json()['converted']==31
    assert client.post('/api/targets/convert-crs',json={'crs':'WGS84'}).json()['converted']==31
    roundtrip=client.get('/api/targets').json()
    for target in roundtrip:
        assert (target['lon'],target['lat'])==pytest.approx(expected[target['id']],abs=4e-6)


def test_target_conversion_does_not_change_existing_job_snapshot(client):
    ident,_=imported(client)
    created=job(client,ident,target_ids=[1],times=['08:00'])
    snapshot=created['config']['targets'][0].copy()
    assert client.post('/api/targets/convert-crs',json={'crs':'WGS84'}).status_code==200
    current=client.get('/api/jobs/'+created['id']).json()['config']['targets'][0]
    assert current==snapshot


def test_bulk_delete_jobs_and_cascade_results(client):
    ident,_=imported(client)
    first=run(client,job(client,ident,target_ids=[1],times=['08:00']))
    second=job(client,ident,target_ids=[1],times=['08:00'])
    blocked=client.post('/api/jobs/bulk-delete',json={'ids':[first['id'],second['id']]})
    assert blocked.status_code==409 and '先暂停' in blocked.json()['detail']
    assert client.post(f"/api/jobs/{second['id']}/pause").status_code==200
    response=client.post('/api/jobs/bulk-delete',json={'ids':[first['id'],second['id']]})
    assert response.status_code==200 and response.json()['deleted']==2
    assert client.get('/api/jobs').json()==[]
    assert client.app.state.db.one('SELECT count(*) n FROM results')['n']==0
    assert client.post('/api/jobs/bulk-delete',json={'ids':[first['id']]}).status_code==404


def test_bulk_delete_datasets_files_and_usage_guard(client):
    used,_=imported(client)
    free1,_=imported(client)
    free2,_=imported(client)
    used_file=next((client.app.state.db.path.parent/'uploads').glob(used+'.*'))
    free_files=[next((client.app.state.db.path.parent/'uploads').glob(i+'.*')) for i in (free1,free2)]
    created=job(client,used,target_ids=[1],times=['08:00'])
    datasets={d['id']:d for d in client.get('/api/datasets').json()}
    assert datasets[used]['job_count']==1 and datasets[free1]['job_count']==0
    blocked=client.post('/api/datasets/bulk-delete',json={'ids':[used,free1]})
    assert blocked.status_code==409 and '相关任务' in blocked.json()['detail']
    assert all(path.exists() for path in free_files) and used_file.exists()
    response=client.post('/api/datasets/bulk-delete',json={'ids':[free1,free2]})
    assert response.status_code==200 and response.json()['deleted']==2
    assert not any(path.exists() for path in free_files)
    assert client.app.state.db.one('SELECT count(*) n FROM points WHERE dataset_id IN (?,?)',(free1,free2))['n']==0
    client.post(f"/api/jobs/{created['id']}/pause")
    client.post('/api/jobs/bulk-delete',json={'ids':[created['id']]})
    assert client.post('/api/datasets/bulk-delete',json={'ids':[used]}).status_code==200
    assert not used_file.exists()

def test_immediate_uses_actual_time_without_date(client):
    ident,_=imported(client)
    r=client.post('/api/jobs',json={'name':'立即采集','dataset_id':ident,'target_ids':[1]})
    assert r.status_code==200
    j=r.json()
    assert j['config']['collection_mode']=='immediate' and len(j['config']['departures'])==1
    done=run(client,j)
    assert done['collection']['started_at'] and done['collection']['finished_at']
    assert done['collection']['requested_count']==13
    again=job(client,ident,times=['08:00','12:00','18:00'])
    assert len(again['config']['departures'])==1 and again['total']==39
    assert 'time_scenario' not in again['config']

def test_scheduled_wait_windows_and_no_cache(client,monkeypatch):
    from datetime import datetime
    ident,_=imported(client)
    j=job(client,ident,date='2099-09-15',times=['08:00','14:00'],collection_mode='scheduled',window_minutes=20)
    worker=client.app.state.worker
    assert j['status']=='scheduled' and worker.ready_job() is None
    due=datetime.fromisoformat(j['config']['departures'][0]).timestamp()
    monkeypatch.setattr('services.worker.time.time',lambda:due+1)
    assert worker.ready_job()['id']==j['id']
    first=run(client,j)
    assert first['status']=='scheduled' and first['success']==39
    rows=client.get('/api/jobs/'+j['id']+'/results').json()['rows']
    assert all(r['requested_at'] and not r['cache_hit'] for r in rows if r['status']=='success')
    assert client.app.state.db.one('SELECT count(*) n FROM cache')['n']==0
    assert worker.ready_job() is None
    second=datetime.fromisoformat(j['config']['departures'][1]).timestamp()
    monkeypatch.setattr('services.worker.time.time',lambda:second+1201)
    done=run(client,first)
    assert done['status']=='partial' and done['failed']==39 and done['success']==39
    failed=client.get('/api/jobs/'+j['id']+'/results?status=failed').json()['rows']
    assert all('错过采集窗口' in r['error'] and r['requested_at'] is None for r in failed)

def test_schedule_past_rejected_and_pause(client):
    ident,_=imported(client)
    body=dict(name='定时',dataset_id=ident,target_ids=[1],date='2000-01-01',times=['08:00'],collection_mode='scheduled')
    assert client.post('/api/jobs',json=body).status_code==400
    j=job(client,ident,date='2099-09-15',times=['08:00'],collection_mode='scheduled')
    assert client.delete('/api/jobs/'+j['id']).status_code==409
    assert client.post('/api/jobs/'+j['id']+'/pause').json()['status']=='paused'
    assert client.app.state.worker.ready_job() is None
    assert client.post('/api/jobs/'+j['id']+'/resume').status_code==200
    assert client.app.state.worker.ready_job() is None

def test_api_build_contract(client):
    r=client.get('/api/version')
    assert r.json()['version']=='2026.09.21.10' and r.headers['cache-control']=='no-store'
    for endpoint in ('targets/convert-crs','jobs/bulk-delete','datasets/bulk-delete'):
        assert client.post('/api/'+endpoint,json={}).status_code==422
    assert client.get('/static/app.js').headers['cache-control']=='no-store'

def research_fixture(client):
    from datetime import datetime,timedelta
    ident,_=imported(client,content='ID,名称,lon,lat\na,甲,114,30\nb,乙,115,31\n'.encode())
    j=job(client,ident,target_ids=[1],dates=['2099-09-15','2099-09-16'],date=None,times=['08:00','14:00'],collection_mode='scheduled',window_minutes=20,day_type='workday',random_seed=42)
    db=client.app.state.db;db.execute("UPDATE jobs SET status='running' WHERE id=?",(j['id'],))
    client.app.state.worker.prepare(j,j['config'])
    rows=db.rows('SELECT * FROM results WHERE job_id=? ORDER BY departure,point_pk',(j['id'],))
    for i,r in enumerate(rows):
        minutes=[10,20,20,30,12,22,22,32][i]
        start=datetime.fromisoformat(r['departure'])+timedelta(seconds=5+i)
        db.execute("UPDATE results SET status='success',duration_seconds=?,requested_at=?,last_requested_at=?,fetched_at=?,attempts=1 WHERE id=?",(minutes*60,start.isoformat(),start.isoformat(),(start+timedelta(seconds=1)).isoformat(),r['id']))
    db.execute("UPDATE jobs SET status='completed',success=8 WHERE id=?",(j['id'],))
    return j,rows

def test_multiday_creation_and_order(client):
    from services.worker import order_key
    ident,_=imported(client)
    j=job(client,ident,dates=['2099-09-16','2099-09-15','2099-09-15'],times=['08:00','14:00'],collection_mode='scheduled',random_seed=123)
    assert len(j['config']['departures'])==4 and j['total']==156
    db=client.app.state.db;db.execute("UPDATE jobs SET status='running' WHERE id=?",(j['id'],));client.app.state.worker.prepare(j,j['config'])
    rows=db.rows('SELECT * FROM results WHERE job_id=? ORDER BY departure,query_order,id',(j['id'],))
    assert len(rows)==156 and all(r['query_order']==order_key(j['config'],r['departure'],r['point_pk'],r['target_id']) for r in rows)
    first=[(r['point_pk'],r['target_id']) for r in rows[:39]]
    second=[(r['point_pk'],r['target_id']) for r in rows[39:78]]
    assert set(first)==set(second) and first!=second
    original=[r['query_order'] for r in rows]
    client.app.state.worker.prepare(j,j['config'])
    assert [r['query_order'] for r in db.rows('SELECT query_order FROM results WHERE job_id=? ORDER BY departure,query_order,id',(j['id'],))]==original
    body=dict(name='窗口检查',dataset_id=ident,target_ids=[1],dates=['2099-09-15'],times=['08:00','08:10'],collection_mode='scheduled',window_minutes=20)
    assert client.post('/api/jobs',json=body).status_code==400
    assert client.post('/api/jobs',json={**body,'times':['08:00'],'dates':['2000-01-01']}).status_code==400

def test_research_matched_statistics_and_export(client):
    import math
    j,rows=research_fixture(client)
    r=client.get('/api/jobs/'+j['id']+'/research')
    assert r.status_code==200,r.text
    report=r.json()
    assert len(report['windows'])==4 and all(w['completeness_pct']==100 and w['eligible']==2 for w in report['windows'])
    stats=report['statistics']
    assert [s['mean'] for s in stats]==[16,26]
    assert all(s['common_od']==2 and s['valid_days']==2 and s['sd']==pytest.approx(math.sqrt(2)) for s in stats)
    assert report['paired'][0]['mean']==10 and report['paired'][0]['paired_days']==2
    out=client.get('/api/jobs/'+j['id']+'/export?kind=research')
    assert out.status_code==200
    wb=load_workbook(io.BytesIO(out.content))
    assert wb.sheetnames==['采集完整性','多日统计','时段配对比较','方法与参数']
    assert wb['采集完整性'].max_row==5
    assert report['random_seed']==42

def test_research_missing_and_invalid_success(client):
    j,rows=research_fixture(client);db=client.app.state.db
    db.execute("UPDATE results SET status='failed',error='错过采集窗口，未请求',requested_at=NULL,last_requested_at=NULL WHERE id=?",(rows[-1]['id'],))
    report=client.get('/api/jobs/'+j['id']+'/research').json()
    assert report['windows'][-1]['missed']==1 and report['windows'][-1]['completeness_pct']==50
    assert [s['mean'] for s in report['statistics']]==[11,21]
    assert all(s['common_od']==1 for s in report['statistics'])
    db.execute('UPDATE results SET cache_hit=1 WHERE id=?',(rows[-2]['id'],))
    report=client.get('/api/jobs/'+j['id']+'/research').json()
    assert report['windows'][-1]['excluded_success']==1
    assert all(s['common_od']==0 and s['mean'] is None for s in report['statistics'])

def test_research_unprepared_windows(client):
    from services.research import research_report
    from datetime import datetime,timezone
    ident,_=imported(client)
    j=job(client,ident,dates=['2099-09-15','2099-09-16'],times=['08:00'],collection_mode='scheduled')
    future=research_report(client.app.state.db,j['id'])
    assert all(w['pending']==39 and w['missed']==0 and w['completeness_pct']==0 for w in future['windows'])
    past=research_report(client.app.state.db,j['id'],at=datetime(2100,1,1,tzinfo=timezone.utc))
    assert all(w['pending']==39 and w['missed']==39 and w['failed']==0 for w in past['windows'])
    assert past['statistics'][0]['mean'] is None
    assert client.get('/api/jobs/'+j['id']+'/export?kind=research').status_code==200

def test_research_outside_window_is_excluded(client):
    from datetime import datetime,timedelta
    j,rows=research_fixture(client);db=client.app.state.db
    r=rows[0]
    late=(datetime.fromisoformat(r['departure'])+timedelta(minutes=20)).isoformat()
    db.execute('UPDATE results SET last_requested_at=? WHERE id=?',(late,r['id']))
    report=client.get('/api/jobs/'+j['id']+'/research').json()
    assert report['windows'][0]['success']==2 and report['windows'][0]['eligible']==1
    assert report['statistics'][0]['common_od']==1

def test_key_status_does_not_block_home(client,monkeypatch):
    import threading
    from services.settings import Settings
    gate=threading.Event()
    settings=Settings(client.app.state.db,client.app.state.db.path.parent)
    def waiting(*args):
        gate.wait(2)
        return ''
    monkeypatch.setattr('services.settings.keyring.get_password',waiting)
    try:
        start=time.monotonic()
        assert settings.public()['has_key'] is None
        assert time.monotonic()-start<.7
    finally: gate.set()

def test_bulk_delete_targets_atomic_and_snapshot(client):
    dataset,_=imported(client)
    created=job(client,dataset)
    before=client.get('/api/jobs/'+created['id']).json()['config']['targets']
    assert client.post('/api/targets/bulk-delete',json={'ids':[]}).status_code==422
    assert client.post('/api/targets/bulk-delete',json={'ids':[1,999999]}).status_code==409
    assert any(t['id']==1 for t in client.get('/api/targets').json())
    response=client.post('/api/targets/bulk-delete',json={'ids':[1,2,2]})
    assert response.status_code==200 and response.json()['deleted']==2
    assert not {1,2}&{t['id'] for t in client.get('/api/targets').json()}
    assert client.get('/api/jobs/'+created['id']).json()['config']['targets']==before
    assert client.get('/api/datasets').json()
    remaining=[t['id'] for t in client.get('/api/targets').json()]
    assert client.post('/api/targets/bulk-delete',json={'ids':remaining}).json()['deleted']==len(remaining)
    assert client.get('/api/targets').json()==[]

@pytest.mark.parametrize('kind',['csv','xlsx'])
def test_target_import_flow(client,kind):
    before=client.get('/api/targets').json()
    ident,stats=imported(client,kind)
    assert stats['valid']>0 and stats['invalid']>0
    points=client.get(f'/api/datasets/{ident}/points?size=200').json()['rows']
    assert client.post(f'/api/datasets/{ident}/as-targets').status_code==200
    after=client.get('/api/targets').json()
    assert after[:len(before)]==before
    added=after[len(before):]
    assert len(added)==stats['valid']
    expected=[(p['name'] or p['point_id'],p['lon'],p['lat'],p['crs']) for p in points if p['valid']]
    assert [(p['name'],p['lon'],p['lat'],p['crs']) for p in added]==expected
    provenance=client.app.state.db.rows('SELECT target_id,dataset_id,row_no,raw FROM target_sources WHERE dataset_id=? ORDER BY row_no',(ident,))
    assert len(provenance)==stats['valid'] and all(row['raw'] for row in provenance)

def test_destination_import_detects_crs_and_keeps_raw_row(client):
    body='destination_id,destination_name,final_route_lon,final_route_lat,coordinate_system,route_poi_name,ready_for_routing\nD001,游客中心,112.72,29.67,GCJ-02,入口,TRUE\n'
    uploaded=client.post('/api/imports',files={'file':('destinations.csv',body.encode())})
    assert uploaded.status_code==200 and uploaded.json()['crs']=='GCJ-02'
    data=uploaded.json()
    checked=client.post(f"/api/imports/{data['id']}/validate",json={'mapping':data['mapping'],'crs':data['crs']})
    assert checked.status_code==200 and checked.json()['valid']==1
    assert client.post(f"/api/datasets/{data['id']}/as-targets").status_code==200
    target=client.app.state.db.one('''SELECT t.target_id,s.dataset_id,s.row_no,s.raw
                                      FROM target_sources s JOIN targets t ON t.id=s.target_id
                                      WHERE s.dataset_id=?''',(data['id'],))
    assert target and target['target_id']=='D001' and '游客中心' in target['raw']

def test_translated_research_headers_are_auto_detected(client):
    """The delivered Chinese-header research files need no manual mapping."""
    source_body = (
        '客源地编号,省份,客源市场单元,城市人口（GHSL）,客源地经度（GCJ-02）,客源地纬度（GCJ-02）,路线查询坐标系,是否可用于路线查询\n'
        'O001,湖北省,武汉市,1000000,114.3055,30.5928,GCJ-02,是\n'
    ).encode('utf-8-sig')
    source = client.post('/api/imports', files={'file': ('31客源单元_最终数据.csv', source_body)})
    assert source.status_code == 200, source.text
    source_data = source.json()
    assert source_data['mapping']['point_id'] == 0
    assert source_data['mapping']['lon'] == 4
    assert source_data['mapping']['lat'] == 5
    assert source_data['mapping']['coordinate_system'] == 6
    assert source_data['mapping']['population'] == 3
    checked = client.post(
        f"/api/imports/{source_data['id']}/validate",
        json={'mapping': source_data['mapping'], 'crs': source_data['crs']},
    )
    assert checked.status_code == 200 and checked.json()['valid'] == 1

    destination_body = (
        '重点村编号,省份,城市,县区,村名,最终路线经度（GCJ-02）,最终路线纬度（GCJ-02）,路线查询坐标系,是否可用于路线查询\n'
        'D001,湖北省,武汉市,黄陂区,示例村,114.45,30.88,GCJ-02,是\n'
    ).encode('utf-8-sig')
    destination = client.post('/api/imports', files={'file': ('94重点村_最终数据.csv', destination_body)})
    assert destination.status_code == 200, destination.text
    destination_data = destination.json()
    assert destination_data['mapping']['point_id'] == 0
    assert destination_data['mapping']['name'] == 4
    assert destination_data['mapping']['lon'] == 5
    assert destination_data['mapping']['lat'] == 6
    assert destination_data['mapping']['coordinate_system'] == 7
    checked = client.post(
        f"/api/imports/{destination_data['id']}/validate",
        json={'mapping': destination_data['mapping'], 'crs': destination_data['crs']},
    )
    assert checked.status_code == 200 and checked.json()['valid'] == 1
    added = client.post(f"/api/datasets/{destination_data['id']}/as-targets")
    assert added.status_code == 200 and added.json()['added'] == 1

def test_standard_targets_and_legacy_summary(client):
    from services.exporter import summaries
    ident,_=imported(client)
    db=client.app.state.db
    db.execute("UPDATE targets SET institution_id='H1',institution_name='医院' WHERE id IN (1,2)")
    assert client.post('/api/targets/classify',json={'ids':[1,2]}).status_code in (404,405)
    j=run(client,job(client,ident,target_ids=[1,2]))
    assert 'institution_summary' not in j['config']
    assert list(summaries(db,j['id'],j['config']['departures'][0]))[0]['success_targets']==2
    cfg=j['config'];cfg['institution_summary']=True
    db.execute('UPDATE jobs SET config=? WHERE id=?',(json.dumps(cfg),j['id']))
    assert list(summaries(db,j['id'],cfg['departures'][0]))[0]['success_targets']==1

def test_capacity_limits_and_acknowledgement(client,monkeypatch):
    from services.settings import Settings
    monkeypatch.setattr(Settings,'key',lambda _: 'test')
    ident,_=imported(client)
    body=dict(name='预估',dataset_id=ident,target_ids=[1,2,3],provider='amap',collection_mode='scheduled',dates=['2099-01-01','2099-01-02'],times=['08:00','14:00'],window_minutes=1,estimate_seconds=10,acknowledge_estimate=True)
    r=client.post('/api/jobs/estimate',json=body).json()
    assert r['routes_per_batch']==39 and r['batches']==4 and r['total_routes']==156
    assert r['planning_minutes']==pytest.approx(2.1666666667) and r['window_exceeded'] and r['warnings']
    assert client.post('/api/jobs',json=body).status_code==400
    body['acknowledge_capacity']=True
    assert client.post('/api/jobs',json=body).status_code==200
    body['provider']='mock'
    assert client.post('/api/jobs/estimate',json=body).json()['planning_minutes'] is None
